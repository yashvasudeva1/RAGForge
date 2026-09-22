"""
DOCX (Microsoft Word) document loader using python-docx.
"""

from __future__ import annotations

from pathlib import Path
from typing import TYPE_CHECKING

from app.core.constants import FileType
from app.core.exceptions import FileReadError
from app.core.logging import get_logger
from app.domain.documents.models import Document, DocumentMetadata
from app.ingestion.loaders.base import BaseLoader

logger = get_logger(__name__)


class DocxLoader(BaseLoader):
    """
    Load Microsoft Word (.docx) files.

    Parameters
    ----------
    include_headers:
        Include heading paragraphs in extracted text.
    include_tables:
        Convert tables to plain text and include them.
    """

    def __init__(
        self,
        include_headers: bool = True,
        include_tables: bool = True,
    ) -> None:
        self.include_headers = include_headers
        self.include_tables = include_tables

    def load(self, source: str) -> list[Document]:
        try:
            from docx import Document as DocxDocument  # type: ignore[import-untyped]
        except ImportError as exc:
            raise ImportError("python-docx is required. Run: pip install python-docx") from exc

        path = Path(source)
        if not path.exists():
            raise FileReadError(source, FileNotFoundError(f"File not found: {source}"))

        try:
            doc = DocxDocument(str(path))
            parts: list[str] = []

            for para in doc.paragraphs:
                if not para.text.strip():
                    continue
                if not self.include_headers and para.style.name.startswith("Heading"):
                    continue
                parts.append(para.text)

            if self.include_tables:
                for table in doc.tables:
                    rows = []
                    for row in table.rows:
                        rows.append(" | ".join(cell.text.strip() for cell in row.cells))
                    parts.append("\n".join(rows))

            content = "\n\n".join(parts)

            # Extract core properties
            props = doc.core_properties
            document = Document(
                content=content,
                metadata=DocumentMetadata(
                    source=str(path),
                    file_type=FileType.DOCX,
                    file_name=path.name,
                    file_size_bytes=path.stat().st_size,
                    title=props.title or None,
                    author=props.author or None,
                ),
            )
            logger.debug("Loaded DOCX: %s (%d chars)", path.name, len(content))
            return [document]

        except Exception as exc:
            raise FileReadError(str(path), exc) from exc
