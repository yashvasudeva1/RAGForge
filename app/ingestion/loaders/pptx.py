"""
PowerPoint (.pptx) document loader using python-pptx.
"""

from __future__ import annotations

from pathlib import Path

from app.core.constants import FileType
from app.core.exceptions import FileReadError
from app.core.logging import get_logger
from app.domain.documents.models import Document, DocumentMetadata
from app.ingestion.loaders.base import BaseLoader

logger = get_logger(__name__)


class PPTXLoader(BaseLoader):
    """
    Load PowerPoint .pptx files.

    Parameters
    ----------
    include_notes:
        Include speaker notes text.
    slide_per_document:
        When True, each slide becomes a separate Document.
    """

    def __init__(
        self,
        include_notes: bool = True,
        slide_per_document: bool = True,
    ) -> None:
        self.include_notes = include_notes
        self.slide_per_document = slide_per_document

    def load(self, source: str) -> list[Document]:
        try:
            from pptx import Presentation  # type: ignore[import-untyped]
        except ImportError as exc:
            raise ImportError("python-pptx is required. Run: pip install python-pptx") from exc

        path = Path(source)
        if not path.exists():
            raise FileReadError(source, FileNotFoundError(f"File not found: {source}"))

        try:
            prs = Presentation(str(path))
            slide_texts: list[str] = []

            for slide in prs.slides:
                parts: list[str] = []
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for para in shape.text_frame.paragraphs:
                            text = para.text.strip()
                            if text:
                                parts.append(text)
                if self.include_notes and slide.has_notes_slide:
                    notes = slide.notes_slide.notes_text_frame.text.strip()
                    if notes:
                        parts.append(f"[Notes] {notes}")
                slide_texts.append("\n".join(parts))

            if self.slide_per_document:
                documents = [
                    Document(
                        content=text,
                        metadata=DocumentMetadata(
                            source=str(path),
                            file_type=FileType.PPTX,
                            file_name=path.name,
                            page_count=len(prs.slides),
                            extra={"slide_number": i + 1},
                        ),
                    )
                    for i, text in enumerate(slide_texts)
                    if text.strip()
                ]
            else:
                full_content = "\n\n".join(t for t in slide_texts if t.strip())
                documents = [
                    Document(
                        content=full_content,
                        metadata=DocumentMetadata(
                            source=str(path),
                            file_type=FileType.PPTX,
                            file_name=path.name,
                            page_count=len(prs.slides),
                            file_size_bytes=path.stat().st_size,
                        ),
                    )
                ]

            logger.debug("Loaded PPTX: %s (%d slides)", path.name, len(prs.slides))
            return documents

        except Exception as exc:
            raise FileReadError(str(path), exc) from exc
